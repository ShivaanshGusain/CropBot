"""The frozen gold set converts to jsonl without changing anything (kg.eval.gold)."""

import json

import pytest
from openpyxl import Workbook, load_workbook

from kg.eval.gold import (CHANGELOG, FACT_SHEETS, GOLD_DIR, WORKBOOK, GoldError, check, out_name, problems,
                          read_workbook, render)


@pytest.fixture(scope="module")
def book():
    return read_workbook()


def test_every_cell_is_kept_as_stored(book):
    """Read independently with openpyxl: each non-empty cell is in its record, same type and value."""
    wb = load_workbook(WORKBOOK, data_only=False)
    for ws in wb.worksheets:
        records = {r["_row"]: r for r in book["sheets"][ws.title]["records"]}
        header = book["sheets"][ws.title]["header"]
        for row in ws.iter_rows(min_row=2):
            if all(c.value is None for c in row):
                assert row[0].row not in records
                continue
            rec = records[row[0].row]
            for c in row:
                assert type(rec[header[c.column - 1]]) is type(c.value) and rec[header[c.column - 1]] == c.value


def test_jsonl_round_trips_the_records(book):
    files = render(book)
    for name, s in book["sheets"].items():
        lines = files[out_name(name)].splitlines()
        assert [json.loads(line) for line in lines] == s["records"]
        assert {json.loads(line)["_sheet"] for line in lines} == {name}
    assert [s for s in book["sheets"]] == [*FACT_SHEETS, CHANGELOG]
    log = book["sheets"][CHANGELOG]["records"][0]
    assert (log["_sheet"], log["sheet"], log["fact_id"]) == ("Changelog", "A", "A38")   # its own "sheet" column kept


def test_committed_files_equal_a_fresh_conversion(book):
    assert check(book) == []


def test_another_workbook_is_refused(tmp_path):
    with pytest.raises(GoldError, match="never edited here"):
        read_workbook(WORKBOOK, sha256="0" * 64)


def test_formulas_are_refused(tmp_path):
    wb = Workbook()
    wb.active.append(["fact_id", "amount"])
    wb.active.append(["X1", "=1+1"])
    path = tmp_path / "f.xlsx"
    wb.save(path)
    with pytest.raises(GoldError, match="formula"):
        read_workbook(path, sha256=None)


def test_problems_are_reported_not_fixed(book):
    p = problems(book)
    assert p["rows"] == {"A": 47, "B": 35, "C": 19, "Changelog": 128}
    assert p["duplicate_fact_ids"] == [] and not p["headers_differ"]
    shifted = ["A25", "A26", "A27", "A28", "A30", "A37", "A44", "B37", "B42", "C7", "C8"]
    assert p["possible_column_shift"] == shifted
    c7 = next(r for r in book["sheets"]["C"]["records"] if r["fact_id"] == "C7")
    assert c7["quote"] is None and c7["other_values"].startswith("When 90 % of the panicle")   # as stored
    manifest = json.loads((GOLD_DIR / "gold_manifest.json").read_text(encoding="utf-8"))
    assert manifest["sha256"] == book["sha256"]
